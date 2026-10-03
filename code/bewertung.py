"""Jedem Artikel einen Status geben: sicher / prüfen / manuell – mit Begründung."""
from pathlib import Path

SICHER, PRUEFEN, MANUELL = "sicher", "prüfen", "manuell"


def lade_korrekturen(ordner):
    """Liest alle ausgefüllten Prüflisten (*.xlsx) aus dem Ordner. Spätere Dateien überschreiben frühere."""
    from openpyxl import load_workbook

    korrekturen = {}
    for datei in sorted(Path(ordner).glob("*.xlsx")):
        if datei.name.startswith("~$"):  # Excel-Sperrdatei
            continue
        wb = load_workbook(datei, read_only=True, data_only=True)
        for ws in wb.worksheets:
            zeilen = ws.iter_rows(values_only=True)
            kopf = [str(k or "") for k in next(zeilen, [])]
            if "Artikelnummer" not in kopf or "Freigabe (ja)" not in kopf:
                continue
            i = {k: kopf.index(k) for k in kopf}
            for z in zeilen:
                werte = {
                    "name": z[i["Name (korrigiert)"]],
                    "beschreibung": z[i["Beschreibung (korrigiert)"]],
                    "bild_url": z[i["Bild-URL (korrigiert)"]],
                    "freigabe": z[i["Freigabe (ja)"]],
                    "hersteller": z[i["Hersteller (korrigiert)"]] if "Hersteller (korrigiert)" in i else None,
                    "notiz": z[i["Notiz"]] if "Notiz" in i else None,
                }
                if z[i["Artikelnummer"]] and any(v for k, v in werte.items() if k != "notiz"):
                    korrekturen[str(z[i["Artikelnummer"]])] = {k: str(v or "").strip() for k, v in werte.items()}
        wb.close()
    return korrekturen


def bewerten(artikel, cache, korrekturen, gesperrt=None):
    """Ergänzt jeden Artikel um produktdaten, status und gruende.

    gesperrt: {artikelnummer: grund} aus der Stichprobe – solche Artikel gehen zurück auf "prüfen",
    bis sie per Korrektur erneut freigegeben werden.
    """
    gesperrt = gesperrt or {}
    for a in artikel:
        daten = cache.get(a) or {}
        gruende = []
        status = SICHER

        if not a["hersteller"]:
            gruende.append(f"Hersteller zum Kürzel '{a['kuerzel']}' unbekannt")
            status = MANUELL
        elif not a["hersteller_gesichert"]:
            gruende.append("Hersteller-Zuordnung nur vermutet")
            status = PRUEFEN
        for g in a["bestellnummer_gruende"]:
            gruende.append(g)
            status = MANUELL
        if not a["quelle"] and status != MANUELL:
            gruende.append("keine automatische Quelle für diesen Hersteller")
            status = MANUELL
        elif a["quelle"] and not daten:
            gruende.append("noch nicht abgefragt")
            status = MANUELL
        elif daten.get("nicht_gefunden"):
            gruende.append("auf Herstellerseite nicht gefunden")
            status = MANUELL
        elif daten:
            if daten.get("hinweise"):
                gruende += daten["hinweise"]
                if status == SICHER:
                    status = PRUEFEN
            if daten.get("gefunden_mit") and daten["gefunden_mit"] != a["bestellnummer"]:
                gruende.append(f"gefunden über Artikelnummer ({daten['gefunden_mit']}), nicht Bestellnummer")
                if status == SICHER:
                    status = PRUEFEN
        if a["variante"] and status == SICHER:
            gruende.append("Lagervariante (Endung in Artikelnummer)")
            status = PRUEFEN

        k = korrekturen.get(a["artikelnummer"])
        if k and k.get("hersteller"):
            a["hersteller"] = k["hersteller"]
        if k:
            daten = {**daten, **{f: k[f] for f in ("name", "beschreibung", "bild_url") if k.get(f)}}
            if k.get("freigabe", "").strip().lower() in ("ja", "x", "1"):
                status, gruende = SICHER, ["manuell geprüft"]
            elif daten.get("name"):  # Vorschlag aus Recherche/Korrektur, aber noch nicht freigegeben
                status = PRUEFEN
                gruende = ["Vorschlag aus Recherche: " + (k.get("notiz") or "bitte prüfen")]

        if a["artikelnummer"] in gesperrt and not (k and k.get("freigabe", "").strip().lower() in ("ja", "x", "1")):
            status, gruende = PRUEFEN, [gesperrt[a["artikelnummer"]]]

        a["produkt"] = {
            "name": daten.get("name", ""),
            "beschreibung": daten.get("beschreibung", ""),
            "bild_url": daten.get("bild_url", ""),
            "quelle_url": daten.get("quelle_url", ""),
        }
        a["status"] = status
        a["gruende"] = gruende
    return artikel
